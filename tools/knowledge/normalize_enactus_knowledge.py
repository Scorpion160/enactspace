#!/usr/bin/env python3
"""
Enactus ESP knowledge normalizer v2.

Changes vs v1:
- Handles source files literally named ".pdf" (Pathlib normally reports no suffix).
- Runs pdftotext against a temporary ASCII-only path to avoid Windows/MiKTeX
  failures on emoji/decomposed-Unicode source paths.
- Compares discovered source count with catalog/sources.csv.
- Never modifies source files.
- No OCR.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import defaultdict
from pathlib import Path

try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None
try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None
try:
    import xlrd
except Exception:
    xlrd = None
try:
    from pptx import Presentation
except Exception:
    Presentation = None
try:
    from odf.opendocument import load as odf_load
    from odf import text as odf_text
    from odf.teletype import extractText as odf_extract_text
except Exception:
    odf_load = odf_text = odf_extract_text = None
try:
    from striprtf.striprtf import rtf_to_text
except Exception:
    rtf_to_text = None
try:
    from PIL import Image
except Exception:
    Image = None

TEXT_EXTS = {".txt", ".md"}
CSV_EXTS = {".csv"}
JSON_EXTS = {".json"}
PDF_EXTS = {".pdf"}
DOCX_EXTS = {".docx"}
DOC_EXTS = {".doc"}
XLSX_EXTS = {".xlsx"}
XLS_EXTS = {".xls"}
PPTX_EXTS = {".pptx"}
PPT_EXTS = {".ppt"}
ODF_EXTS = {".odt", ".ods", ".odp"}
RTF_EXTS = {".rtf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".webp"}

SUPPORTED_EXTS = (
    TEXT_EXTS | CSV_EXTS | JSON_EXTS | PDF_EXTS | DOCX_EXTS | DOC_EXTS |
    XLSX_EXTS | XLS_EXTS | PPTX_EXTS | PPT_EXTS | ODF_EXTS | RTF_EXTS |
    IMAGE_EXTS
)

EMAIL_RE = re.compile(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", re.I)
PHONE_RE = re.compile(
    r"(?<!\d)(?:\+?221[\s.\-]?)?(?:7[05678])(?:[\s.\-]?\d{2,3}){2,4}(?!\d)"
)
PRIVACY_TERMS_RE = re.compile(
    r"\b(?:CNI|carte d'identit[eé]|passeport|adresse|t[eé]l[eé]phone|contact|"
    r"autorisation parentale|signature|renvoi|disciplin(?:e|aire)|absence(?:s)?|"
    r"date de naissance|naissance|matricule)\b",
    re.I,
)


def logical_ext(path: Path) -> str:
    """Return the semantic extension, including weird hidden names like '.pdf'."""
    ext = path.suffix.lower()
    if ext:
        return ext
    hidden = path.name.lower()
    if hidden in SUPPORTED_EXTS:
        return hidden
    return ""


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while chunk := f.read(chunk_size):
            h.update(chunk)
    return h.hexdigest().upper()


def rel_posix(path: Path, root: Path) -> str:
    return path.relative_to(root).as_posix()


def safe_scalar(value) -> str:
    if value is None:
        return ""
    return str(value).replace("\r", " ").replace("\n", " ").strip()


def read_text_guess(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("utf-8", errors="replace")


def run_pdftotext(path: Path, exe: str) -> tuple[str, str]:
    # MiKTeX's Windows pdftotext can fail on emoji / decomposed Unicode paths.
    # Copy to an ASCII-only temporary filename first; never modify the source.
    with tempfile.TemporaryDirectory(prefix="enactus_pdf_") as td:
        temp_pdf = Path(td) / "input.pdf"
        shutil.copyfile(path, temp_pdf)
        proc = subprocess.run(
            [exe, "-layout", "-enc", "UTF-8", str(temp_pdf), "-"],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=120,
        )
    if proc.returncode != 0:
        raise RuntimeError(proc.stderr.strip() or f"pdftotext exit={proc.returncode}")
    return proc.stdout.strip(), "pdftotext-temp-ascii"


def extract_docx(path: Path) -> tuple[str, str]:
    if DocxDocument is None:
        raise RuntimeError("python-docx not installed")
    doc = DocxDocument(str(path))
    out: list[str] = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if "heading" in style or "titre" in style:
            level = 2
            m = re.search(r"(\d+)", style)
            if m:
                level = min(6, max(2, int(m.group(1)) + 1))
            out.append(f"{'#' * level} {t}")
        else:
            out.append(t)
    for ti, table in enumerate(doc.tables, start=1):
        out.append(f"\n## Tableau {ti}")
        for ri, row in enumerate(table.rows, start=1):
            vals = [safe_scalar(cell.text) for cell in row.cells]
            while vals and vals[-1] == "":
                vals.pop()
            if vals:
                out.append(f"- Ligne {ri}: " + " | ".join(vals))
    return "\n\n".join(out).strip(), "python-docx"


def extract_xlsx(path: Path) -> tuple[str, str]:
    if load_workbook is None:
        raise RuntimeError("openpyxl not installed")
    wb = load_workbook(str(path), read_only=True, data_only=False)
    out: list[str] = []
    try:
        for ws in wb.worksheets:
            out.append(f"# Feuille: {ws.title}")
            emitted = 0
            for ri, row in enumerate(ws.iter_rows(values_only=True), start=1):
                vals = [safe_scalar(v) for v in row]
                while vals and vals[-1] == "":
                    vals.pop()
                if vals and any(v != "" for v in vals):
                    out.append(f"- Ligne {ri}: " + " | ".join(vals))
                    emitted += 1
            if emitted == 0:
                out.append("_Feuille vide ou sans valeurs lisibles._")
    finally:
        wb.close()
    return "\n".join(out).strip(), "openpyxl"


def extract_xls(path: Path) -> tuple[str, str]:
    if xlrd is None:
        raise RuntimeError("xlrd not installed")
    wb = xlrd.open_workbook(str(path), on_demand=True)
    out: list[str] = []
    try:
        for sheet in wb.sheets():
            out.append(f"# Feuille: {sheet.name}")
            for ri in range(sheet.nrows):
                vals = [safe_scalar(sheet.cell_value(ri, ci)) for ci in range(sheet.ncols)]
                while vals and vals[-1] == "":
                    vals.pop()
                if vals and any(v != "" for v in vals):
                    out.append(f"- Ligne {ri + 1}: " + " | ".join(vals))
    finally:
        wb.release_resources()
    return "\n".join(out).strip(), "xlrd"


def extract_pptx(path: Path) -> tuple[str, str]:
    if Presentation is None:
        raise RuntimeError("python-pptx not installed")
    prs = Presentation(str(path))
    out: list[str] = []
    for si, slide in enumerate(prs.slides, start=1):
        out.append(f"# Diapositive {si}")
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                txt = (shape.text or "").strip()
                if txt:
                    out.append(txt)
            if getattr(shape, "has_table", False):
                for ri, row in enumerate(shape.table.rows, start=1):
                    vals = [safe_scalar(cell.text) for cell in row.cells]
                    while vals and vals[-1] == "":
                        vals.pop()
                    if vals:
                        out.append(f"- Tableau ligne {ri}: " + " | ".join(vals))
    return "\n\n".join(out).strip(), "python-pptx"


def extract_odf(path: Path) -> tuple[str, str]:
    if odf_load is None or odf_text is None or odf_extract_text is None:
        raise RuntimeError("odfpy not installed")
    doc = odf_load(str(path))
    parts: list[str] = []
    for node in doc.getElementsByType(odf_text.H):
        t = odf_extract_text(node).strip()
        if t:
            parts.append(f"## {t}")
    for node in doc.getElementsByType(odf_text.P):
        t = odf_extract_text(node).strip()
        if t:
            parts.append(t)
    return "\n\n".join(parts).strip(), "odfpy"


def extract_rtf(path: Path) -> tuple[str, str]:
    if rtf_to_text is None:
        raise RuntimeError("striprtf not installed")
    return rtf_to_text(read_text_guess(path)).strip(), "striprtf"


def extract_csv(path: Path) -> tuple[str, str]:
    text = read_text_guess(path)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    out: list[str] = []
    for ri, row in enumerate(csv.reader(text.splitlines(), dialect), start=1):
        vals = [safe_scalar(v) for v in row]
        while vals and vals[-1] == "":
            vals.pop()
        if vals:
            out.append(f"- Ligne {ri}: " + " | ".join(vals))
    return "\n".join(out).strip(), "csv"


def extract_json(path: Path) -> tuple[str, str]:
    obj = json.loads(read_text_guess(path))
    return "```json\n" + json.dumps(obj, ensure_ascii=False, indent=2) + "\n```", "json"


def extract_image_metadata(path: Path) -> tuple[str, str]:
    if Image is None:
        raise RuntimeError("Pillow not installed")
    with Image.open(path) as img:
        body = (
            "## Métadonnées image\n\n"
            f"- Format: {img.format}\n"
            f"- Dimensions: {img.width} × {img.height}\n"
            f"- Mode: {img.mode}\n"
        )
    return body, "Pillow-metadata"


def extract_content(path: Path, pdftotext_exe: str) -> tuple[str, str, str]:
    ext = logical_ext(path)
    if ext in TEXT_EXTS:
        return read_text_guess(path).strip(), "text", "OK"
    if ext in CSV_EXTS:
        txt, ex = extract_csv(path)
        return txt, ex, "OK"
    if ext in JSON_EXTS:
        txt, ex = extract_json(path)
        return txt, ex, "OK"
    if ext in PDF_EXTS:
        txt, ex = run_pdftotext(path, pdftotext_exe)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT_LAYER")
    if ext in DOCX_EXTS:
        txt, ex = extract_docx(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in DOC_EXTS:
        return "", "none", "UNSUPPORTED_LEGACY_DOC"
    if ext in XLSX_EXTS:
        txt, ex = extract_xlsx(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in XLS_EXTS:
        txt, ex = extract_xls(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in PPTX_EXTS:
        txt, ex = extract_pptx(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in PPT_EXTS:
        return "", "none", "UNSUPPORTED_LEGACY_PPT"
    if ext in ODF_EXTS:
        txt, ex = extract_odf(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in RTF_EXTS:
        txt, ex = extract_rtf(path)
        return txt, ex, ("OK" if re.search(r"\w", txt) else "NO_TEXT")
    if ext in IMAGE_EXTS:
        txt, ex = extract_image_metadata(path)
        return txt, ex, "METADATA_ONLY"
    return "", "none", "UNSUPPORTED"


def privacy_signals(text: str) -> list[str]:
    signals: list[str] = []
    if EMAIL_RE.search(text):
        signals.append("EMAIL")
    if PHONE_RE.search(text):
        signals.append("PHONE")
    if PRIVACY_TERMS_RE.search(text):
        signals.append("PRIVACY_TERM")
    return signals


def write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def read_catalog_count(catalog_csv: Path) -> tuple[int, list[str]]:
    if not catalog_csv.exists():
        return 0, []
    with catalog_csv.open("r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    return len(rows), [r.get("RelativePath", "").replace("\\", "/") for r in rows]


def resolve_pdftotext(requested: str | None) -> str | None:
    return shutil.which(requested or "pdftotext")


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo", default=".")
    ap.add_argument(
        "--pdftotext",
        default=None,
    )
    args = ap.parse_args()

    repo = Path(args.repo).resolve()
    sources = repo / "docs/enactus_knowledge/sources"
    normalized = repo / "docs/enactus_knowledge/normalized"
    catalog = repo / "docs/enactus_knowledge/catalog"
    pdftotext_exe = resolve_pdftotext(args.pdftotext)

    if not sources.exists():
        print(f"ERROR: sources not found: {sources}", file=sys.stderr)
        return 2
    if pdftotext_exe is None:
        requested = args.pdftotext or "pdftotext"
        print(f"ERROR: pdftotext not found or not executable: {requested}", file=sys.stderr)
        return 3

    normalized.mkdir(parents=True, exist_ok=True)
    catalog.mkdir(parents=True, exist_ok=True)

    files = sorted(
        p for p in sources.rglob("*")
        if p.is_file() and p.name != ".gitattributes" and logical_ext(p) in SUPPORTED_EXTS
    )

    catalog_count, catalog_paths = read_catalog_count(catalog / "sources.csv")
    discovered_paths = {rel_posix(p, sources) for p in files}
    missing_from_discovery = sorted(set(catalog_paths) - discovered_paths)

    print(f"[INFO] Source files discovered: {len(files)}")
    print(f"[INFO] catalog/sources.csv rows: {catalog_count}")
    if missing_from_discovery:
        print("[WARN] Catalog entries not discovered:")
        for item in missing_from_discovery:
            print(f"  - {item}")

    sha_to_paths: dict[str, list[Path]] = defaultdict(list)
    file_sha: dict[Path, str] = {}

    for idx, path in enumerate(files, start=1):
        sha = sha256_file(path)
        file_sha[path] = sha
        sha_to_paths[sha].append(path)
        if idx % 50 == 0 or idx == len(files):
            print(f"[HASH] {idx}/{len(files)}")

    duplicate_rows: list[dict] = []
    for sha, paths in sorted(sha_to_paths.items()):
        if len(paths) > 1:
            duplicate_rows.append({
                "SHA256": sha,
                "Count": len(paths),
                "Paths": " | ".join(rel_posix(p, sources) for p in paths),
            })
    write_csv(catalog / "duplicates_exact.csv", ["SHA256", "Count", "Paths"], duplicate_rows)

    result_by_sha: dict[str, dict] = {}
    failure_rows: list[dict] = []
    privacy_rows: list[dict] = []
    unique_shas = sorted(sha_to_paths)

    for idx, sha in enumerate(unique_shas, start=1):
        primary = sha_to_paths[sha][0]
        source_rel = rel_posix(primary, sources)
        out_rel = Path("by_sha") / sha[:2] / f"{sha}.md"
        out_path = normalized / out_rel
        out_path.parent.mkdir(parents=True, exist_ok=True)

        try:
            body, extractor, status = extract_content(primary, pdftotext_exe)
            error = ""
        except Exception as exc:
            body = ""
            extractor = "none"
            status = "ERROR"
            error = f"{type(exc).__name__}: {exc}"

        frontmatter = (
            "---\n"
            f"source_path: {json.dumps(source_rel, ensure_ascii=False)}\n"
            f"source_sha256: {json.dumps(sha)}\n"
            f"extension: {json.dumps(logical_ext(primary))}\n"
            'classification: "INTERNAL_KNOWLEDGE"\n'
            f"extraction_status: {json.dumps(status)}\n"
            f"extractor: {json.dumps(extractor)}\n"
            "---\n\n"
            f"# {primary.name}\n\n"
        )

        if status == "NO_TEXT_LAYER":
            rendered = "_Aucune couche texte exploitable détectée. OCR non exécuté._\n"
        elif status == "NO_TEXT":
            rendered = "_Aucun contenu textuel exploitable détecté._\n"
        elif status.startswith("UNSUPPORTED"):
            rendered = "_Format historique non normalisé automatiquement._\n"
        elif status == "ERROR":
            rendered = "_Échec d'extraction. Voir `catalog/extraction_failures.csv`._\n"
        elif status == "METADATA_ONLY":
            rendered = "_Image conservée comme source ; aucun OCR exécuté._\n\n" + body.strip() + "\n"
        else:
            rendered = body.strip() + "\n"

        out_path.write_text(frontmatter + rendered, encoding="utf-8")

        result = {
            "SHA256": sha,
            "NormalizedPath": rel_posix(out_path, repo),
            "Status": status,
            "Extractor": extractor,
            "Chars": len(body),
            "Error": error,
        }
        result_by_sha[sha] = result

        if error or status in {"NO_TEXT_LAYER", "NO_TEXT"} or status.startswith("UNSUPPORTED"):
            failure_rows.append({"SourcePath": source_rel, **result})

        signals = privacy_signals(body)
        if signals:
            privacy_rows.append({
                "SourcePath": source_rel,
                "SHA256": sha,
                "NormalizedPath": result["NormalizedPath"],
                "Signals": "|".join(signals),
            })

        if idx % 25 == 0 or idx == len(unique_shas):
            print(f"[EXTRACT] {idx}/{len(unique_shas)} unique files")

    mapping_rows: list[dict] = []
    report_rows: list[dict] = []

    for path in files:
        sha = file_sha[path]
        r = result_by_sha[sha]
        source_rel = rel_posix(path, sources)
        mapping_rows.append({
            "SourcePath": source_rel,
            "SHA256": sha,
            "NormalizedPath": r["NormalizedPath"],
            "IsExactDuplicate": "YES" if len(sha_to_paths[sha]) > 1 else "NO",
        })
        report_rows.append({
            "SourcePath": source_rel,
            "SHA256": sha,
            "Extension": logical_ext(path),
            "Status": r["Status"],
            "Extractor": r["Extractor"],
            "Chars": r["Chars"],
            "NormalizedPath": r["NormalizedPath"],
            "Error": r["Error"],
        })

    write_csv(
        catalog / "source_to_normalized.csv",
        ["SourcePath", "SHA256", "NormalizedPath", "IsExactDuplicate"],
        mapping_rows,
    )
    write_csv(
        catalog / "normalization_report.csv",
        ["SourcePath", "SHA256", "Extension", "Status", "Extractor", "Chars", "NormalizedPath", "Error"],
        report_rows,
    )
    write_csv(
        catalog / "extraction_failures.csv",
        ["SourcePath", "SHA256", "NormalizedPath", "Status", "Extractor", "Chars", "Error"],
        failure_rows,
    )
    write_csv(
        catalog / "privacy_review_candidates.csv",
        ["SourcePath", "SHA256", "NormalizedPath", "Signals"],
        privacy_rows,
    )

    status_counts: dict[str, int] = defaultdict(int)
    for row in report_rows:
        status_counts[row["Status"]] += 1

    print("\n========== NORMALIZATION SUMMARY ==========")
    print(f"Source files         : {len(files)}")
    print(f"Catalog rows         : {catalog_count}")
    print(f"Unique SHA-256       : {len(unique_shas)}")
    print(f"Exact duplicate sets : {len(duplicate_rows)}")
    for key in sorted(status_counts):
        print(f"{key:21}: {status_counts[key]}")
    print(f"Privacy candidates   : {len(privacy_rows)}")
    print(f"Normalized directory : {normalized}")
    print(f"Catalog directory    : {catalog}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
